"""
Extrait les tables et les formules du classeur de référence vers donnees/*.csv.

Usage :
    python outils/extraire_excel.py <classeur_formules.xlsx> [classeur_valeurs.xlsx] [dossier_sortie]

Le classeur des valeurs doit contenir des valeurs calculées (fichier enregistré par Excel
ou recalculé par LibreOffice) ; par défaut, c'est le même fichier.
Chaque ligne produite garde la cellule Excel d'où elle vient (colonne « source »).
"""
import csv, os, re, sys, warnings

import openpyxl
from openpyxl.utils import column_index_from_string as col_num, get_column_letter as col_lettre

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from formules_excel import ADRESSES_SAISIE, NOMS, FormuleNonGeree, analyser, en_python, nombre, simplifier
from seuils_re2020.variables import ALTITUDES, ZONES

warnings.filterwarnings("ignore")

# ─── Paramètres ──────────────────────────────────────────────────────────────
F_FORMULES = sys.argv[1]
F_VALEURS = sys.argv[2] if len(sys.argv) > 2 else F_FORMULES
SORTIE = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(__file__), "..", "donnees")
SEP = ";"                                  # séparateur CSV (lisible dans Excel FR)
# Résultats DH réutilisés par le DH Effinergie : désignent le DH_max de l'usage courant
REFS_DH_MAX = {"Calcul_dh!C32", "Calcul_dh!C35", "Calcul_dh!C43", "Calcul_dh!C50"}

# Blocs « HLOOKUP("ICI", …) » : (coefficient, feuille, ligne ICI, ligne calcul, colonnes)
BLOCS_ICI = [
    ("Mbsurf_moy", "Calcul_bbio", 51, 52, "G", "M"),
    ("Mbsurf_tot", "Calcul_bbio", 58, 59, "G", "AJ"),
    ("MbHSP", "Calcul_bbio", 99, 100, "E", "J"),
    ("Mcsurf_moy", "Calcul_energie", 84, 85, "G", "AC"),
    ("Mcsurf_tot", "Calcul_energie", 96, 97, "G", "AP"),
    ("McHSP", "Calcul_energie", 116, 117, "E", "J"),
    ("Misurf_moy", "Calcul_matiere", 42, 43, "G", "Z"),
    ("Misurf_tot", "Calcul_matiere", 48, 50, "G", "AE"),
    ("Mided", "Calcul_matiere", 69, 70, "E", "AR"),
    ("Miagrement_ext", "Calcul_matiere", 78, 79, "E", "H"),
    ("MiHSP", "Calcul_matiere", 85, 86, "E", "I"),
]
# Formules isolées traduites telles quelles
FORMULES_SIMPLES = [
    ("Mbcombles", "Calcul_bbio", "E46"),
    ("Mccombles", "Calcul_energie", "E79"),
    ("Micombles", "Calcul_matiere", "E37"),
    ("Mipv", "Calcul_matiere", "E62"),
    ("Miclim_RCU", "Calcul_matiere", "E88"),
    ("Mcgeo_exception_MI", "Calcul_energie", "D75"),
    ("DH_effinergie_Q8", "Calcul_dh", "Q8"),
]

wf = openpyxl.load_workbook(F_FORMULES)
wv = openpyxl.load_workbook(F_VALEURS, data_only=True)


# ─── Outils ──────────────────────────────────────────────────────────────────
def norm(texte):
    """Normalisation des libellés : comparaison Excel insensible à la casse."""
    return str(texte).replace("\xa0", " ").strip().lower()


def val(feuille, adresse):
    return wv[feuille][adresse].value


def formule_brute(feuille, adresse):
    v = wf[feuille][adresse].value
    return getattr(v, "text", v)


# noms définis -> cellule
NOM_PAR_CELLULE = {}
for nom, defn in wf.defined_names.items():
    cible = defn.attr_text.replace("$", "")
    if "!" in cible and ":" not in cible and "#REF" not in cible:
        NOM_PAR_CELLULE[cible.replace("'", "")] = nom


def resolveur(feuille):
    def resoudre(ref):
        if "!" in ref:
            f, a = ref.rsplit("!", 1)
            f = f.strip("'")
        else:
            f, a = feuille, ref
        a = a.replace("$", "")
        if ":" in a:
            raise FormuleNonGeree(f"plage {ref}")
        cle = f"{f}!{a}"
        if cle in ADRESSES_SAISIE:
            return ADRESSES_SAISIE[cle]
        if cle in REFS_DH_MAX:
            return "dh_max"
        if cle in NOM_PAR_CELLULE and NOM_PAR_CELLULE[cle].lower() in NOMS:
            return NOMS[NOM_PAR_CELLULE[cle].lower()]
        try:
            texte = traduire(f, a)
            # parenthèses conservées autour d'une sous-expression insérée
            return texte if re.fullmatch(r"[\w.]+|'[^']*'|-?\d+(\.\d+)?", texte) else f"({texte})"
        except FormuleNonGeree:
            # cellule de libellé calculée par une recherche : on prend sa valeur (journalisée)
            v = val(f, a)
            REPLIS.append(f"{cle} = {v!r}  ({formule_brute(f, a)})")
            return nombre(v) if isinstance(v, (int, float)) else repr(norm(v))
    return resoudre


REPLIS = []   # cellules remplacées par leur valeur, à relire


def traduire(feuille, adresse):
    """Contenu d'une cellule -> expression Python."""
    brut = formule_brute(feuille, adresse)
    if isinstance(brut, str) and brut.startswith("="):
        return simplifier(en_python(analyser(brut), resolveur(feuille)))
    if brut is None:
        return "0"
    if isinstance(brut, (int, float)):
        return nombre(brut)
    return repr(norm(brut))


def ecrire(nom, entetes, lignes):
    os.makedirs(SORTIE, exist_ok=True)
    with open(os.path.join(SORTIE, nom), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=SEP)
        w.writerow(entetes)
        w.writerows(lignes)
    print(f"{nom:32} {len(lignes):5} lignes")


# ─── Usages ──────────────────────────────────────────────────────────────────
E = "Enumération"
PRINCIPAUX = {norm(val(E, f"C{r}")): val(E, f"C{r}") for r in range(4, 22)}
USAGES = []   # (id, principal, détaillé)
for i, r in enumerate(range(4, 32), start=1):
    detail = val(E, f"E{r}")
    principal = detail.split(" - ", 1)[0] if " - " in detail else detail
    assert norm(principal) in PRINCIPAUX, detail
    USAGES.append((i, PRINCIPAUX[norm(principal)], detail))
PAR_DETAIL = {norm(d): i for i, _, d in USAGES}
PAR_PRINCIPAL = {}
for i, p, _ in USAGES:
    PAR_PRINCIPAL.setdefault(norm(p), []).append(i)
CE3 = {norm(val(E, f"Q{r}")) for r in range(3, 12) if val(E, f"Q{r}")}
ecrire("usages.csv", ["id", "usage_principal", "usage_detaille", "categorie3_possible"],
       [(i, p, d, int(norm(p) in CE3 or norm(d) in CE3)) for i, p, d in USAGES])


def ids_usage(libelle, detaille):
    cle = norm(libelle)
    if detaille:
        return [PAR_DETAIL[cle]]
    return PAR_PRINCIPAL[cle]


# ─── Infobulles : colonne I de la feuille Input_resultats ────────────────────
# ligne Excel -> champ de saisie ; plusieurs lignes peuvent viser le même champ
AIDES_LIGNES = {9: "altitude", 10: "sref", 11: "scombles", 12: "nb_logements", 13: "nb_logements",
                14: "exposition_bruit", 19: "rcu_classe", 22: "categorie", 23: "climatise",
                26: "sagrement", 27: "hsp", 28: "pc_apres_juillet_2026", 29: "classe_inertie"}


def texte_aide(brut):
    """Texte d'une cellule d'aide : pour une formule SI(…), la branche « vrai » ; URL -> lien."""
    if isinstance(brut, str) and brut.startswith("="):
        chaines = re.findall(r'"((?:[^"]|"")*)"', brut)
        brut = chaines[0].replace('""', '"') if chaines else ""
    texte = str(brut).strip()
    return re.sub(r"(https?://\S+)", r"[\1](\1)", texte)


aides = {}
for ligne, champ in AIDES_LIGNES.items():
    t = texte_aide(formule_brute("Input_resultats", f"I{ligne}"))
    if t:
        texte_prec, src_prec = aides.get(champ, ("", ""))
        aides[champ] = ((texte_prec + ". " + t if texte_prec else t), (src_prec + ";" if src_prec else "") + f"Input_resultats!I{ligne}")
ecrire("aides.csv", ["champ", "texte", "source"], [(c, t, s) for c, (t, s) in aides.items()])

# ─── Zones climatiques ───────────────────────────────────────────────────────
Z = "zone climatique"
lignes = []
for r in range(3, wv[Z].max_row + 1):
    num, dep, zone = val(Z, f"B{r}"), val(Z, f"C{r}"), val(Z, f"D{r}")
    if dep and zone in ZONES:
        lignes.append((str(num).zfill(2) if isinstance(num, int) else num, dep, zone, f"'{Z}'!D{r}"))
ecrire("zones_departements.csv", ["numero", "departement", "zone", "source"], lignes)

# ─── Valeurs moyennes ────────────────────────────────────────────────────────
lignes = []
for i, _, d in USAGES:
    rb = next(r for r in range(5, 33) if norm(val("Calcul_bbio", f"D{r}")) == norm(d))
    re_ = next(r for r in range(6, 34) if norm(val("Calcul_energie", f"D{r}")) == norm(d))
    lignes.append((i, val("Calcul_bbio", f"E{rb}"), val("Calcul_energie", f"E{re_}"),
                   val("Calcul_energie", f"F{re_}"), f"Calcul_bbio!E{rb};Calcul_energie!E{re_}:F{re_}"))
ecrire("moyennes.csv", ["usage_id", "bbio_maxmoyen", "cepnr_maxmoyen", "cep_maxmoyen", "source"], lignes)

# Icconstruction_maxmoyen (valeurs officielles ; la dérogation IGH est appliquée par le moteur)
lignes = []
for r in range(5, 23):
    principal = val("Calcul_matiere", f"D{r}")
    for col, periode in (("E", "2022 à 2024"), ("F", "2025 à 2027"), ("L", "2028 à 2030"), ("M", "à partir de 2031")):
        v = val("Calcul_matiere", f"{col}{r}")
        for i in PAR_PRINCIPAL[norm(principal)]:
            lignes.append((i, periode, v if isinstance(v, (int, float)) else "NA", f"Calcul_matiere!{col}{r}"))
ecrire("icconstruction_moyen.csv", ["usage_id", "periode", "valeur", "source"], lignes)

# Icénergie_maxmoyen : clé « usage détaillé - RCU / - autres cas »
lignes = []
for r in range(6, 62):
    cle = val("Calcul_energie", f"K{r}")
    if not cle:
        continue
    detail, cas = cle.rsplit(" - ", 1)
    for col, periode in (("L", "2022 à 2024"), ("M", "2025 à 2027"), ("N", "à partir de 2028")):
        v = val("Calcul_energie", f"{col}{r}")
        lignes.append((PAR_DETAIL[norm(detail)], "oui" if cas == "RCU" else "non", periode,
                       v if isinstance(v, (int, float)) else "NA", f"Calcul_energie!{col}{r}"))
ecrire("icenergie_moyen.csv", ["usage_id", "rcu", "periode", "valeur", "source"], lignes)


# ─── Tables géographiques (Mbgéo, Mcgéo, Migéo) ─────────────────────────────
def grille(feuille, l_zone0, col_decalage, l_decalage0, nb, cle_principale=False):
    """Lit une table zone × (usage, altitude) indexée par un décalage de colonne."""
    sortie = []
    for k in range(nb):
        r = l_decalage0 + k
        lib = val(feuille, f"D{r}")
        dec = val(feuille, f"{col_decalage}{r}")
        ids = PAR_PRINCIPAL[norm(lib)] if cle_principale else [PAR_DETAIL[norm(lib)]]
        for a, alt in enumerate(ALTITUDES):
            col = col_lettre(col_num("E") + dec + a)
            for z in range(8):
                adr = f"{col}{l_zone0 + z}"
                for i in ids:
                    sortie.append((i, alt, val(feuille, f"D{l_zone0 + z}"), val(feuille, adr), f"{feuille}!{adr}"))
    return sortie


ecrire("mbgeo.csv", ["usage_id", "altitude", "zone", "valeur", "source"], grille("Calcul_bbio", 37, "F", 5, 28))
ecrire("mcgeo.csv", ["usage_id", "altitude", "zone", "valeur", "source"], grille("Calcul_energie", 66, "G", 6, 28))
ecrire("migeo.csv", ["usage_id", "altitude", "zone", "valeur", "source"],
       grille("Calcul_matiere", 28, "I", 5, 18, cle_principale=True))
# exception Mcgéo maisons 2022-2024 (valeurs avant le décret 2024-1258)
ecrire("mcgeo_exception.csv", ["zone", "valeur", "source"],
       [(val("Calcul_energie", f"D{r}"), val("Calcul_energie", f"E{r}"), f"Calcul_energie!E{r}") for r in (76, 77)])

# ─── Mccat : zone × catégorie ────────────────────────────────────────────────
lignes = []
for r in range(6, 34):
    i = PAR_DETAIL[norm(val("Calcul_energie", f"D{r}"))]
    dec = val("Calcul_energie", f"I{r}")
    for cat in (1, 2, 3):
        col_n = col_num("E") + dec + cat - 1
        for z in range(8):
            adr = f"{col_lettre(col_n)}{102 + z}"
            v = val("Calcul_energie", adr) if col_n <= col_num("BC") else None
            lignes.append((i, cat, ZONES[z], v if isinstance(v, (int, float)) else "NA", f"Calcul_energie!{adr}"))
ecrire("mccat.csv", ["usage_id", "categorie", "zone", "valeur", "source"], lignes)

# ─── Mbbruit : colonnes BR1 / BR2-BR3 / Cat 3, avec renvoi « X » vers une table par zone ──
lignes = []
for r in range(64, 92):
    i = PAR_DETAIL[norm(val("Calcul_bbio", f"D{r}"))]
    for col, cle in (("E", "BR1"), ("F", "BR2_BR3"), ("G", "CAT3")):
        v = val("Calcul_bbio", f"{col}{r}")
        if v == "X":
            k = val("Calcul_bbio", f"H{r}")
            colz = col_lettre(col_num("I") + k - 1)
            for z in range(8):
                lignes.append((i, cle, val("Calcul_bbio", f"I{64 + z}"), val("Calcul_bbio", f"{colz}{64 + z}"),
                               f"Calcul_bbio!{colz}{64 + z}"))
        else:
            lignes.append((i, cle, "*", v if isinstance(v, (int, float)) else "NA", f"Calcul_bbio!{col}{r}"))
ecrire("mbbruit.csv", ["usage_id", "exposition", "zone", "valeur", "source"], lignes)

# ─── Seuils Miinfra / Mivrd par usage principal ──────────────────────────────
lignes = []
for k in range(18):
    col = col_lettre(col_num("E") + k)
    for i in PAR_PRINCIPAL[norm(val("Calcul_matiere", f"{col}53"))]:
        lignes.append((i, val("Calcul_matiere", f"{col}54"), val("Calcul_matiere", f"{col}59"),
                       f"Calcul_matiere!{col}54;{col}59"))
ecrire("seuils_lots.csv", ["usage_id", "seuil_infra", "seuil_vrd", "source"], lignes)


# ─── Blocs « ICI » : règles conditionnelles ──────────────────────────────────
def termes_et(arbre):
    if arbre[0] == "fn" and arbre[1] == "AND":
        return [t for a in arbre[2] for t in termes_et(a)]
    if arbre[0] == "op" and arbre[1] == "*":
        return termes_et(arbre[2]) + termes_et(arbre[3])
    return [arbre]


def terme_usage(t, feuille):
    """Reconnaît « usage = cellule » ; renvoie (libellé, détaillé ?) ou None."""
    if t[0] == "op" and t[1] == "=":
        g, d = t[2], t[3]
        for a, b in ((g, d), (d, g)):
            nom = a[1].lower() if a[0] == "nom" else ADRESSES_SAISIE.get(a[1].replace("$", "")) if a[0] == "ref" else None
            if nom in ("usage", "usage_detaillé", "usage_detaille") and b[0] == "ref":
                ref = b[1].replace("$", "")
                f, adr = ref.rsplit("!", 1) if "!" in ref else (feuille, ref)
                return val(f.strip("'"), adr), nom != "usage"
    return None


regles = []
for coef, feuille, l_ici, l_calc, c0, c1 in BLOCS_ICI:
    for n in range(col_num(c0), col_num(c1) + 1):
        col = col_lettre(n)
        ici = formule_brute(feuille, f"{col}{l_ici}")
        if not isinstance(ici, str) or not ici.startswith("="):
            continue
        formule = traduire(feuille, f"{col}{l_calc}")
        if "MATCH(\"ICI\"" in ici or "COUNTIF(" in ici or re.search(rf'[A-Z]+{l_ici}="NON"', ici):
            regles.append((coef, "*", "defaut", formule, f"{feuille}!{col}{l_calc}"))
            continue
        arbre = analyser(ici)
        assert arbre[0] == "fn" and arbre[1] == "IF", ici
        usages, conditions = None, []
        for t in termes_et(arbre[2][0]):
            if t[0] == "ref" and formule_brute(feuille, t[1].replace("$", "").split("!")[-1]) is None:
                continue                      # ET() d'Excel ignore les cellules vides
            u = terme_usage(t, feuille)
            if u:
                usages = ids_usage(*u)
            else:
                terme = simplifier(en_python(t, resolveur(feuille)))
                conditions.append(terme if re.fullmatch(r"[\w.]+", terme) else f"({terme})")
        assert usages, (coef, col, ici)
        regles.append((coef, ",".join(map(str, usages)), simplifier(" and ".join(conditions)) if len(conditions) == 1 else " and ".join(conditions) or "True",
                       formule, f"{feuille}!{col}{l_calc}"))

for coef, feuille, adr in FORMULES_SIMPLES:
    regles.append((coef, "*", "True", traduire(feuille, adr), f"{feuille}!{adr}"))

# ─── DH_max : un bloc par usage dans Calcul_dh ───────────────────────────────
for r in range(3, 30):
    lib = val("Calcul_dh", f"B{r}")
    cible = formule_brute("Calcul_dh", f"C{r}").lstrip("=")
    bloc = formule_brute("Calcul_dh", cible)
    usages = PAR_PRINCIPAL.get(norm(lib)) or [PAR_DETAIL[norm(lib)]]
    u = ",".join(map(str, usages))
    if bloc.startswith('=HLOOKUP("ICI"'):
        rng = re.search(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", bloc)
        ca, la, cb, lb = rng.group(1), int(rng.group(2)), rng.group(3), int(rng.group(4))
        for n in range(col_num(ca), col_num(cb) + 1):
            col = col_lettre(n)
            ici = formule_brute("Calcul_dh", f"{col}{la}")
            if not ici:
                continue
            cond = simplifier(en_python(analyser(ici).__getitem__(2)[0], resolveur("Calcul_dh")))
            regles.append(("DH_max", u, cond, traduire("Calcul_dh", f"{col}{lb}"), f"Calcul_dh!{col}{lb}"))
    elif bloc.startswith("=HLOOKUP(categorie"):          # maisons : catégorie seule
        for col in ("G", "H"):
            regles.append(("DH_max", u, f"categorie == {norm(val('Calcul_dh', f'{col}31'))!r}",
                           traduire("Calcul_dh", f"{col}32"), f"Calcul_dh!{col}32"))
    elif bloc.startswith('=VLOOKUP("ICI",F38:I40'):      # logements collectifs : Smoylgt × colonne
        for lr in (38, 39, 40):
            cl = simplifier(en_python(analyser(formule_brute("Calcul_dh", f"F{lr}")).__getitem__(2)[0],
                                      resolveur("Calcul_dh")))
            for col in ("G", "H", "I"):
                cc = simplifier(en_python(analyser(formule_brute("Calcul_dh", f"{col}37")).__getitem__(2)[0],
                                          resolveur("Calcul_dh")))
                regles.append(("DH_max", u, f"({cl}) and ({cc})", traduire("Calcul_dh", f"{col}{lr}"),
                               f"Calcul_dh!{col}{lr}"))
    else:
        raise FormuleNonGeree(f"bloc DH inconnu {bloc}")

# DH Effinergie (label, hors réglementation) : maisons, logements, bureaux, enseignement
for r in range(3, 7):
    lib = val("Calcul_dh", f"E{r}")
    regles.append(("DH_effinergie", ",".join(map(str, PAR_PRINCIPAL[norm(lib)])), "True",
                   traduire("Calcul_dh", f"F{r}"), f"Calcul_dh!F{r}"))

ecrire("regles.csv", ["coefficient", "usages", "condition", "formule", "source"], regles)
print("\nCellules remplacées par leur valeur :")
print("\n".join(sorted(set(REPLIS))))
